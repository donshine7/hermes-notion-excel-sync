from __future__ import annotations

import io
import tempfile
import unittest
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import Workbook

from notion_excel_sync.adapters.excel import ExcelSnapshotReader, SheetSpec
from notion_excel_sync.adapters.onedrive import DriveVersion
from notion_excel_sync.domain import AnalysisPipeline, build_default_registry
from notion_excel_sync.knowledge.provider import AnalyzerKnowledgeSnapshot
from notion_excel_sync.models import (
    ChangeKind,
    KnowledgeRef,
    ProposalAction,
    ProposedChange,
    SourceRef,
)
from notion_excel_sync.persistence.database import StateDatabase
from notion_excel_sync.security.approval import ApprovalService
from notion_excel_sync.workflow.proposals import ProposalService
from notion_excel_sync.workflow.sync import SyncPreparationService


class FakeOneDrive:
    def __init__(self, content: bytes) -> None:
        self.content = content
        self.version = DriveVersion(
            "2.0",
            datetime(2026, 7, 16, 1, tzinfo=UTC),
        )

    def list_versions(self, drive_id: str, item_id: str):
        del drive_id, item_id
        return [self.version]

    def download_version(self, drive_id: str, item_id: str, version_id: str):
        del drive_id, item_id
        if version_id != self.version.id:
            raise AssertionError(version_id)
        return self.content

    def download_current(self, drive_id: str, item_id: str):
        del drive_id, item_id
        return self.content

    def select_initial_baseline(self, drive_id: str, item_id: str, cutoff):
        del drive_id, item_id, cutoff
        return self.version


class CommandTimeLocalSource(FakeOneDrive):
    def select_initial_baseline(self, drive_id: str, item_id: str, cutoff):
        del drive_id, item_id, cutoff
        raise AssertionError("Local first run must not select a historical cutoff")


class TwoVersionOneDrive(FakeOneDrive):
    def __init__(self, content: bytes) -> None:
        super().__init__(content)
        self.versions = [
            DriveVersion("1.0", datetime(2026, 7, 15, 1, tzinfo=UTC)),
            self.version,
        ]

    def list_versions(self, drive_id: str, item_id: str):
        del drive_id, item_id
        return list(self.versions)

    def download_version(self, drive_id: str, item_id: str, version_id: str):
        del drive_id, item_id
        if version_id not in {"1.0", "2.0"}:
            raise AssertionError(version_id)
        return self.content


class VersionedWorkbookOneDrive(TwoVersionOneDrive):
    def __init__(self, baseline: bytes, current: bytes) -> None:
        super().__init__(current)
        self.contents = {"1.0": baseline, "2.0": current}

    def download_version(self, drive_id: str, item_id: str, version_id: str):
        del drive_id, item_id
        return self.contents[version_id]

    def download_current(self, drive_id: str, item_id: str):
        del drive_id, item_id
        return self.contents["2.0"]


class NeverCalledPipeline:
    def analyze_changes(self, *args, **kwargs):
        raise AssertionError("No analyzer run is needed when only pending items remain")


class FakeKnowledgeProvider:
    def __init__(self, snapshot: AnalyzerKnowledgeSnapshot) -> None:
        self.snapshot = snapshot
        self.calls = 0

    def prepare(self, changes, pipeline):
        del pipeline
        self.calls += 1
        if not changes:
            raise AssertionError("Knowledge retrieval requires changed Excel rows")
        return self.snapshot


def knowledge_ref() -> KnowledgeRef:
    return KnowledgeRef(
        drive_id="wiki-drive",
        item_id="wiki-item",
        version_id="wiki-version",
        content_hash="wiki-document-hash",
        chunk_locator="paragraph:1#chunk:0",
        chunk_hash="wiki-chunk-hash",
        security_classification="internal",
        display_name="정부지원사업 증빙 지침",
        extractor_version="extractor-v1",
        policy_version="policy-v1",
        effective_from="2026-01-01",
        effective_to="2026-12-31",
        authority="issuing-agency",
        status="verified",
    )


def workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "2026"
    sheet.append(["ID", "상태", "사용자정의필드"])
    sheet.append(["A-1", "착수", "custom"])
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def evidence_workbook_bytes(eligible_amount: int) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "지원사업"
    sheet.append(
        [
            "ID",
            "당소 사건번호",
            "정부지원사업",
            "협약번호",
            "증빙유형",
            "인정건수",
            "인정대상금액",
            "약정총액",
        ]
    )
    sheet.append(
        [
            "E-1",
            "SS-1",
            "2026 수출바우처",
            "VB-2026-01",
            "가출원 1건",
            "1건",
            eligible_amount,
            None,
        ]
    )
    sheet.append(["C-1", "SS-1", None, None, None, None, None, 1_000_000])
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def three_case_evidence_workbook_bytes(first_eligible_amount: int) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "지원사업"
    sheet.append(
        [
            "ID",
            "당소 사건번호",
            "정부지원사업",
            "협약번호",
            "증빙유형",
            "인정건수",
            "인정대상금액",
            "약정총액",
        ]
    )
    for index, eligible_amount in enumerate(
        (first_eligible_amount, 900_000, 900_000),
        start=1,
    ):
        sheet.append(
            [
                f"E-{index}",
                f"SS-{index}",
                "2026 수출바우처",
                "VB-2026-01",
                "가출원 3건",
                "3건",
                eligible_amount,
                None,
            ]
        )
    for index in range(1, 4):
        sheet.append(
            [
                f"C-{index}",
                f"SS-{index}",
                None,
                None,
                None,
                None,
                None,
                300_000,
            ]
        )
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def cross_program_evidence_workbook_bytes(first_eligible_amount: int) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "지원사업"
    sheet.append(
        [
            "ID",
            "당소 사건번호",
            "정부지원사업",
            "협약번호",
            "증빙유형",
            "인정건수",
            "인정대상금액",
            "약정총액",
        ]
    )
    sheet.append(
        [
            "E-A",
            "SS-1",
            "지원사업 A",
            "AGREEMENT-A",
            "가출원 1건",
            "1건",
            first_eligible_amount,
            None,
        ]
    )
    sheet.append(
        [
            "E-B",
            "SS-1",
            "지원사업 B",
            "AGREEMENT-B",
            "가출원 1건",
            "1건",
            500_000,
            None,
        ]
    )
    sheet.append(["C-1", "SS-1", None, None, None, None, None, 600_000])
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


class SyncPreparationTest(unittest.TestCase):
    def test_unconfigured_notion_database_operations_default_to_defer(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            database = StateDatabase(Path(folder) / "state.db")
            database.initialize()
            change = ProposedChange(
                target_database="정부지원사업 증빙",
                entity_key="evidence:1",
                property_name="증빙명",
                kind=ChangeKind.CREATE,
                current_value=None,
                proposed_value="지원사업 증빙안",
                analyzer="government_support_evidence",
                analyzer_version="2.0.0",
                confidence=1.0,
                reason="test",
                source_refs=[SourceRef("drive", "item", "2", "hash", "지원", 2, "A2")],
            )
            proposal = ProposalService(database).create(
                "2",
                "hash",
                "123",
                "chat",
                [change],
                defer_databases={"정부지원사업 증빙"},
            )
            self.assertIs(proposal.operations[0].action, ProposalAction.DEFER)
            self.assertEqual(
                proposal.operations[0].approved_value,
                "지원사업 증빙안",
            )

    def test_unchanged_full_snapshot_cost_is_context_only_for_changed_evidence(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            database = StateDatabase(Path(folder) / "state.db")
            database.initialize()
            database.commit_checkpoint(
                "drive",
                "item",
                "1.0",
                "baseline-hash",
                "previous",
            )
            approval = ApprovalService(
                "a-secure-test-secret-at-least-24-characters",
                {"123"},
                nonce_used=database.nonce_used,
                mark_nonce_used=database.mark_nonce_used,
            )
            service = SyncPreparationService(
                database,
                VersionedWorkbookOneDrive(
                    evidence_workbook_bytes(800_000),
                    evidence_workbook_bytes(900_000),
                ),
                ExcelSnapshotReader([SheetSpec("지원사업", ("ID",))]),
                AnalysisPipeline(build_default_registry()),
                ProposalService(database, approval),
            )

            result = service.prepare(
                drive_id="drive",
                item_id="item",
                initial_cutoff=datetime(2026, 7, 16, tzinfo=UTC),
                requested_by="123",
                chat_id="chat",
            )

            self.assertEqual(result.changes_detected, 1)
            self.assertIsNotNone(result.proposal)
            assert result.proposal is not None
            changes = [operation.change for operation in result.proposal.operations]
            evidence = {
                change.property_name: change
                for change in changes
                if change.target_database == "정부지원사업 증빙"
            }
            self.assertEqual(evidence["실제 비용 합계"].proposed_value, 1_000_000)
            self.assertEqual(len(evidence["관련 실제 비용"].proposed_value), 1)
            self.assertEqual(
                {ref.row for ref in evidence["실제 비용 합계"].source_refs},
                {2, 3},
            )
            self.assertFalse(
                any(change.target_database == "비용·청구" for change in changes)
            )

    def test_one_changed_evidence_row_keeps_the_full_three_case_package(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            database = StateDatabase(Path(folder) / "state.db")
            database.initialize()
            database.commit_checkpoint(
                "drive",
                "item",
                "1.0",
                "baseline-hash",
                "previous",
            )
            approval = ApprovalService(
                "a-secure-test-secret-at-least-24-characters",
                {"123"},
                nonce_used=database.nonce_used,
                mark_nonce_used=database.mark_nonce_used,
            )
            service = SyncPreparationService(
                database,
                VersionedWorkbookOneDrive(
                    three_case_evidence_workbook_bytes(800_000),
                    three_case_evidence_workbook_bytes(900_000),
                ),
                ExcelSnapshotReader([SheetSpec("지원사업", ("ID",))]),
                AnalysisPipeline(build_default_registry()),
                ProposalService(database, approval),
            )

            result = service.prepare(
                drive_id="drive",
                item_id="item",
                initial_cutoff=datetime(2026, 7, 16, tzinfo=UTC),
                requested_by="123",
                chat_id="chat",
            )

            self.assertEqual(result.changes_detected, 1)
            self.assertIsNotNone(result.proposal)
            assert result.proposal is not None
            evidence = {
                operation.change.property_name: operation.change.proposed_value
                for operation in result.proposal.operations
                if operation.change.target_database == "정부지원사업 증빙"
            }
            self.assertEqual(evidence["대상 사건"], ["SS-1", "SS-2", "SS-3"])
            self.assertEqual(evidence["인정 건수"], 3)
            self.assertEqual(evidence["실제 비용 합계"], 900_000)
            self.assertEqual(len(evidence["관련 실제 비용"]), 3)

    def test_changed_package_detects_duplicate_case_in_unchanged_program(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            database = StateDatabase(Path(folder) / "state.db")
            database.initialize()
            database.commit_checkpoint(
                "drive",
                "item",
                "1.0",
                "baseline-hash",
                "previous",
            )
            approval = ApprovalService(
                "a-secure-test-secret-at-least-24-characters",
                {"123"},
                nonce_used=database.nonce_used,
                mark_nonce_used=database.mark_nonce_used,
            )
            service = SyncPreparationService(
                database,
                VersionedWorkbookOneDrive(
                    cross_program_evidence_workbook_bytes(400_000),
                    cross_program_evidence_workbook_bytes(500_000),
                ),
                ExcelSnapshotReader([SheetSpec("지원사업", ("ID",))]),
                AnalysisPipeline(build_default_registry()),
                ProposalService(database, approval),
            )

            result = service.prepare(
                drive_id="drive",
                item_id="item",
                initial_cutoff=datetime(2026, 7, 16, tzinfo=UTC),
                requested_by="123",
                chat_id="chat",
            )

            self.assertIsNotNone(result.proposal)
            assert result.proposal is not None
            duplicates = [
                operation.change.proposed_value
                for operation in result.proposal.operations
                if operation.change.target_database == "정부지원사업 증빙"
                and operation.change.property_name == "중복증빙"
            ]
            self.assertEqual(duplicates, [True])

    def test_shadow_wiki_reports_hits_without_affecting_proposal(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            database = StateDatabase(Path(folder) / "state.db")
            database.initialize()
            database.commit_checkpoint(
                "drive",
                "item",
                "1.0",
                "baseline-hash",
                "previous",
            )
            approval = ApprovalService(
                "a-secure-test-secret-at-least-24-characters",
                {"123"},
                nonce_used=database.nonce_used,
                mark_nonce_used=database.mark_nonce_used,
            )
            knowledge = FakeKnowledgeProvider(
                AnalyzerKnowledgeSnapshot(
                    available=True,
                    binding={
                        "generation_digest": "wiki-shadow-generation",
                        "policy_version": "policy-v1",
                    },
                    shadow_counts={"government_support_evidence": 4},
                )
            )
            service = SyncPreparationService(
                database,
                VersionedWorkbookOneDrive(
                    evidence_workbook_bytes(800_000),
                    evidence_workbook_bytes(900_000),
                ),
                ExcelSnapshotReader([SheetSpec("지원사업", ("ID",))]),
                AnalysisPipeline(build_default_registry()),
                ProposalService(database, approval),
                knowledge_provider=knowledge,
            )

            result = service.prepare(
                drive_id="drive",
                item_id="item",
                initial_cutoff=datetime(2026, 7, 16, tzinfo=UTC),
                requested_by="123",
                chat_id="chat",
            )

            self.assertEqual(knowledge.calls, 1)
            self.assertEqual(result.wiki_generation_digest, "wiki-shadow-generation")
            self.assertEqual(result.wiki_shadow_hits, 4)
            self.assertIsNotNone(result.proposal)
            assert result.proposal is not None
            self.assertEqual(result.proposal.knowledge_binding, {})
            self.assertTrue(
                all(
                    not operation.change.knowledge_refs
                    for operation in result.proposal.operations
                )
            )

    def test_verified_wiki_refs_flow_through_evidence_review_and_history(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            database = StateDatabase(Path(folder) / "state.db")
            database.initialize()
            database.commit_checkpoint(
                "drive",
                "item",
                "1.0",
                "baseline-hash",
                "previous",
            )
            approval = ApprovalService(
                "a-secure-test-secret-at-least-24-characters",
                {"123"},
                nonce_used=database.nonce_used,
                mark_nonce_used=database.mark_nonce_used,
            )
            ref = knowledge_ref()
            binding = {
                "generation_digest": "wiki-verified-generation",
                "policy_version": "policy-v1",
            }
            knowledge = FakeKnowledgeProvider(
                AnalyzerKnowledgeSnapshot(
                    available=True,
                    binding=binding,
                    verified_refs_by_analyzer={
                        "government_support_evidence": (ref,)
                    },
                    shadow_counts={"government_support_evidence": 1},
                )
            )
            service = SyncPreparationService(
                database,
                VersionedWorkbookOneDrive(
                    evidence_workbook_bytes(800_000),
                    evidence_workbook_bytes(1_200_000),
                ),
                ExcelSnapshotReader([SheetSpec("지원사업", ("ID",))]),
                AnalysisPipeline(build_default_registry()),
                ProposalService(database, approval),
                knowledge_provider=knowledge,
            )

            result = service.prepare(
                drive_id="drive",
                item_id="item",
                initial_cutoff=datetime(2026, 7, 16, tzinfo=UTC),
                requested_by="123",
                chat_id="chat",
            )

            self.assertIsNotNone(result.proposal)
            assert result.proposal is not None
            self.assertEqual(result.proposal.knowledge_binding, binding)
            operations = result.proposal.operations
            evidence_operations = [
                operation
                for operation in operations
                if operation.change.target_database == "정부지원사업 증빙"
            ]
            self.assertTrue(evidence_operations)
            self.assertTrue(
                all(operation.change.knowledge_refs == [ref] for operation in evidence_operations)
            )

            review_operations = [
                operation
                for operation in operations
                if operation.change.target_database == "검토함"
                and operation.change.knowledge_refs
            ]
            self.assertTrue(review_operations)
            self.assertTrue(
                all(operation.change.knowledge_refs == [ref] for operation in review_operations)
            )

            evidence_history_keys = {
                f"history:{operation.change.operation_id}"
                for operation in evidence_operations
            }
            evidence_history = [
                operation
                for operation in operations
                if operation.change.target_database == "변경이력"
                and operation.change.entity_key in evidence_history_keys
            ]
            self.assertTrue(evidence_history)
            self.assertTrue(
                all(operation.change.knowledge_refs == [ref] for operation in evidence_history)
            )

    def test_noop_new_version_advances_checkpoint_without_a_proposal(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            database = StateDatabase(Path(folder) / "state.db")
            database.initialize()
            database.commit_checkpoint(
                "drive",
                "item",
                "1.0",
                "old-hash",
                "previous",
            )
            approval = ApprovalService(
                "a-secure-test-secret-at-least-24-characters",
                {"123"},
                nonce_used=database.nonce_used,
                mark_nonce_used=database.mark_nonce_used,
            )
            service = SyncPreparationService(
                database,
                TwoVersionOneDrive(workbook_bytes()),
                ExcelSnapshotReader([SheetSpec("2026", ("ID",))]),
                NeverCalledPipeline(),
                ProposalService(database, approval),
            )
            result = service.prepare(
                drive_id="drive",
                item_id="item",
                initial_cutoff=datetime(2026, 7, 16, tzinfo=UTC),
                requested_by="123",
                chat_id="chat",
            )

            self.assertIsNone(result.proposal)
            self.assertTrue(result.checkpoint_advanced)
            self.assertEqual(
                database.get_checkpoint("drive", "item")["version_id"],
                "2.0",
            )

    def test_local_first_run_reconciles_command_time_snapshot_without_cutoff(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            database = StateDatabase(Path(folder) / "state.db")
            database.initialize()
            approval = ApprovalService(
                "a-secure-test-secret-at-least-24-characters",
                {"123"},
                nonce_used=database.nonce_used,
                mark_nonce_used=database.mark_nonce_used,
            )
            service = SyncPreparationService(
                database,
                CommandTimeLocalSource(workbook_bytes()),
                ExcelSnapshotReader([SheetSpec("2026", ("ID",))]),
                AnalysisPipeline(build_default_registry()),
                ProposalService(database, approval),
                first_run_full_reconcile=True,
            )

            result = service.prepare(
                drive_id="local",
                item_id="local:item",
                initial_cutoff=datetime(1999, 1, 1, tzinfo=UTC),
                requested_by="123",
                chat_id="chat",
            )

            self.assertEqual(result.baseline_version_id, "2.0")
            self.assertEqual(result.current_version_id, "2.0")
            self.assertEqual(result.changes_detected, 1)
            self.assertIsNotNone(result.proposal)
            self.assertIsNone(database.get_checkpoint("local", "local:item"))

    def test_pending_item_reappears_without_a_new_onedrive_version(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            database = StateDatabase(Path(folder) / "state.db")
            database.initialize()
            approval = ApprovalService(
                "a-secure-test-secret-at-least-24-characters",
                {"123"},
                nonce_used=database.nonce_used,
                mark_nonce_used=database.mark_nonce_used,
            )
            proposals = ProposalService(database, approval)
            change = ProposedChange(
                target_database="한국 특허 사건",
                entity_key="A-1",
                property_name="현재상태",
                kind=ChangeKind.UPDATE,
                current_value="접수",
                proposed_value="착수",
                analyzer="workflow_status",
                analyzer_version="1.0.0",
                confidence=1.0,
                reason="Deferred test change",
                source_refs=[
                    SourceRef(
                        "drive",
                        "item",
                        "2.0",
                        "old-hash",
                        "2026",
                        2,
                        "B2",
                        "착수",
                    )
                ],
            )
            queued = proposals.create("2.0", "old-hash", "123", "chat", [change])
            queued.operations[0].action = ProposalAction.DEFER
            database.enqueue_operations(queued)
            database.commit_checkpoint(
                "drive",
                "item",
                "2.0",
                "old-hash",
                "previous",
            )

            service = SyncPreparationService(
                database,
                FakeOneDrive(workbook_bytes()),
                ExcelSnapshotReader([SheetSpec("2026", ("ID",))]),
                NeverCalledPipeline(),
                proposals,
            )
            result = service.prepare(
                drive_id="drive",
                item_id="item",
                initial_cutoff=datetime(2026, 7, 16, tzinfo=UTC),
                requested_by="123",
                chat_id="chat",
            )

            self.assertEqual(result.changes_detected, 0)
            self.assertEqual(result.pending_reintroduced, 1)
            self.assertIsNotNone(result.proposal)
            assert result.proposal is not None
            operation = next(
                item
                for item in result.proposal.operations
                if item.change.target_database == "한국 특허 사건"
            )
            self.assertEqual(operation.change.operation_id, change.operation_id)
            self.assertTrue(result.draft_analyzers)

    def test_wiki_backed_pending_item_is_not_rebound_to_a_new_generation(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            database = StateDatabase(Path(folder) / "state.db")
            database.initialize()
            approval = ApprovalService(
                "a-secure-test-secret-at-least-24-characters",
                {"123"},
                nonce_used=database.nonce_used,
                mark_nonce_used=database.mark_nonce_used,
            )
            proposals = ProposalService(database, approval)
            change = ProposedChange(
                target_database="한국 특허 사건",
                entity_key="A-1",
                property_name="현재상태",
                kind=ChangeKind.UPDATE,
                current_value="접수",
                proposed_value="착수",
                analyzer="workflow_status",
                analyzer_version="1.0.0",
                confidence=1.0,
                reason="Deferred Wiki-backed change",
                source_refs=[
                    SourceRef(
                        "drive",
                        "item",
                        "2.0",
                        "old-hash",
                        "2026",
                        2,
                        "B2",
                        "착수",
                    )
                ],
                knowledge_refs=[knowledge_ref()],
            )
            queued = proposals.create(
                "2.0",
                "old-hash",
                "123",
                "chat",
                [change],
                knowledge_binding={"generation_digest": "original-generation"},
            )
            queued.operations[0].action = ProposalAction.DEFER
            database.enqueue_operations(queued)
            database.commit_checkpoint(
                "drive",
                "item",
                "2.0",
                "old-hash",
                "previous",
            )
            service = SyncPreparationService(
                database,
                FakeOneDrive(workbook_bytes()),
                ExcelSnapshotReader([SheetSpec("2026", ("ID",))]),
                NeverCalledPipeline(),
                proposals,
                knowledge_provider=FakeKnowledgeProvider(
                    AnalyzerKnowledgeSnapshot(
                        available=True,
                        binding={"generation_digest": "current-generation"},
                        verified_refs_by_analyzer={
                            "workflow_status": (knowledge_ref(),)
                        },
                    )
                ),
            )

            with self.assertRaisesRegex(
                RuntimeError,
                "Deferred Wiki-backed operations cannot be rematerialized",
            ):
                service.prepare(
                    drive_id="drive",
                    item_id="item",
                    initial_cutoff=datetime(2026, 7, 16, tzinfo=UTC),
                    requested_by="123",
                    chat_id="chat",
                )

            self.assertEqual(len(database.load_pending_operations()), 1)
            with database.session() as connection:
                audit = connection.execute(
                    "SELECT details FROM audit_log "
                    "WHERE event_type = 'wiki_pending_rematerialization_blocked'"
                ).fetchone()
            self.assertIsNotNone(audit)
            assert audit is not None
            self.assertEqual(audit["details"], '{"operations": 1}')


if __name__ == "__main__":
    unittest.main()
