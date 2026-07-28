from __future__ import annotations

import hashlib
import io
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from notion_excel_sync.models import (
    ChangeKind,
    JsonValue,
    RecordChange,
    SourceRecord,
    SourceRef,
    WorkbookSnapshot,
    canonical_json,
    utc_now,
)


class ExcelReadError(RuntimeError):
    pass


class MissingWorksheetError(ExcelReadError):
    pass


class InvalidHeaderError(ExcelReadError):
    pass


class MissingRecordKeyError(ExcelReadError):
    pass


class DuplicateRecordKeyError(ExcelReadError):
    pass


class SnapshotIdentityMismatch(ValueError):
    pass


@dataclass(frozen=True, slots=True)
class SheetSpec:
    name: str
    key_columns: tuple[str, ...] = ()
    header_row: int = 1
    include_columns: tuple[str, ...] = ()
    required_columns: tuple[str, ...] = ()
    key_groups: tuple[tuple[str, ...], ...] = ()

    def __post_init__(self) -> None:
        if not self.name.strip():
            raise ValueError("sheet name must not be empty")
        if self.header_row < 1:
            raise ValueError("header_row must be at least 1")
        if not self.key_columns and not self.key_groups:
            raise ValueError("key_columns or key_groups are required for stable diffs")
        if self.key_columns and self.key_groups:
            raise ValueError("Use key_columns alternatives or key_groups, not both")
        if len(set(self.key_columns)) != len(self.key_columns):
            raise ValueError("key_columns must not contain duplicates")
        if any(not group for group in self.key_groups):
            raise ValueError("key_groups must not contain an empty group")
        if any(len(set(group)) != len(group) for group in self.key_groups):
            raise ValueError("key_groups must not contain duplicate columns")
        if len(set(self.key_groups)) != len(self.key_groups):
            raise ValueError("key_groups must not contain duplicate groups")

    @property
    def candidate_key_groups(self) -> tuple[tuple[str, ...], ...]:
        """Ordered key strategies; ``key_columns`` are single-column alternatives."""

        if self.key_groups:
            return self.key_groups
        return tuple((column,) for column in self.key_columns)


def _json_value(value: Any) -> JsonValue:
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, Decimal):
        integral = value.to_integral_value()
        return int(integral) if value == integral else float(value)
    return str(value)


def _header_name(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _record_key(
    sheet: str, key_group: tuple[str, ...], values: dict[str, JsonValue]
) -> str:
    payload = {"columns": key_group, "values": [values[column] for column in key_group]}
    return f"{sheet}::{canonical_json(payload)}"


def _resolve_occurrence_keys(records: list[SourceRecord]) -> list[str]:
    """Add deterministic occurrence numbers and report order-sensitive groups."""

    by_base: dict[str, list[SourceRecord]] = {}
    for record in records:
        by_base.setdefault(record.key, []).append(record)
    warnings: list[str] = []
    for base_key, matches in by_base.items():
        full_signatures: dict[str, list[int]] = {}
        for record in matches:
            full_signatures.setdefault(canonical_json(record.values), []).append(
                record.row
            )
        identical_rows = next(
            (rows for rows in full_signatures.values() if len(rows) > 1),
            None,
        )
        if identical_rows:
            locations = ", ".join(str(row) for row in sorted(identical_rows))
            raise DuplicateRecordKeyError(
                f"Indistinguishable duplicate record key {base_key!r} at rows "
                f"{locations}"
            )
        if len(matches) > 1:
            rows = ", ".join(str(record.row) for record in matches)
            warnings.append(
                f"Repeated composite key uses order-sensitive occurrence identity "
                f"at {matches[0].sheet}!{rows}: {base_key}"
            )
        for occurrence, record in enumerate(matches, start=1):
            record.key = f"{base_key}::occurrence::{occurrence}"
    return warnings


class ExcelSnapshotReader:
    """Read selected worksheets from an XLSX payload without modifying the workbook."""

    def __init__(
        self,
        sheet_specs: Iterable[SheetSpec],
        *,
        data_only: bool = True,
        max_workbook_bytes: int = 100 * 1024 * 1024,
        max_cells: int = 2_000_000,
    ) -> None:
        specs = tuple(sheet_specs)
        if not specs:
            raise ValueError("at least one SheetSpec is required")
        names = [spec.name for spec in specs]
        if len(set(names)) != len(names):
            raise ValueError("worksheet specifications must have unique names")
        if max_workbook_bytes <= 0 or max_cells <= 0:
            raise ValueError("workbook limits must be positive")
        self._sheet_specs = specs
        self._data_only = data_only
        self._max_workbook_bytes = max_workbook_bytes
        self._max_cells = max_cells

    def read_file(
        self,
        path: str | Path,
        *,
        drive_id: str,
        item_id: str,
        version_id: str,
    ) -> WorkbookSnapshot:
        """Read a local temporary download.  The file is opened only for binary reading."""

        content = Path(path).read_bytes()
        return self.read_bytes(
            content,
            drive_id=drive_id,
            item_id=item_id,
            version_id=version_id,
        )

    def read_bytes(
        self,
        content: bytes,
        *,
        drive_id: str,
        item_id: str,
        version_id: str,
    ) -> WorkbookSnapshot:
        if len(content) > self._max_workbook_bytes:
            raise ExcelReadError(
                f"Workbook is larger than the configured {self._max_workbook_bytes}-byte limit"
            )
        file_hash = hashlib.sha256(content).hexdigest()
        try:
            workbook = load_workbook(
                io.BytesIO(content),
                read_only=True,
                data_only=self._data_only,
                keep_links=False,
            )
        except Exception as exc:  # openpyxl exposes several format-specific exceptions
            raise ExcelReadError(f"Unable to open XLSX content: {exc}") from exc

        records: list[SourceRecord] = []
        identity_warnings: list[str] = []
        keys_seen: set[str] = set()
        cells_read = 0
        try:
            for spec in self._sheet_specs:
                if spec.name not in workbook.sheetnames:
                    raise MissingWorksheetError(f"Worksheet {spec.name!r} does not exist")
                worksheet = workbook[spec.name]
                header_cells = next(
                    worksheet.iter_rows(min_row=spec.header_row, max_row=spec.header_row),
                    (),
                )
                headers = [_header_name(cell.value) for cell in header_cells]
                while headers and not headers[-1]:
                    headers.pop()
                    header_cells = header_cells[:-1]
                if not headers:
                    raise InvalidHeaderError(
                        f"Worksheet {spec.name!r} has no headers on row {spec.header_row}"
                    )
                named_headers = [header for header in headers if header]
                if len(named_headers) != len(set(named_headers)):
                    duplicates = sorted(
                        {header for header in named_headers if named_headers.count(header) > 1}
                    )
                    raise InvalidHeaderError(
                        f"Worksheet {spec.name!r} has duplicate headers: {duplicates}"
                    )

                available = set(named_headers)
                missing = sorted(set(spec.required_columns) - available)
                if missing:
                    raise InvalidHeaderError(
                        f"Worksheet {spec.name!r} is missing required columns: {missing}"
                    )
                available_key_groups = tuple(
                    group
                    for group in spec.candidate_key_groups
                    if set(group).issubset(available)
                )
                if not available_key_groups:
                    raise InvalidHeaderError(
                        f"Worksheet {spec.name!r} has none of the configured key strategies: "
                        f"{spec.candidate_key_groups}"
                    )
                key_headers = {column for group in available_key_groups for column in group}
                if spec.include_columns:
                    unknown = sorted(set(spec.include_columns) - available)
                    if unknown:
                        raise InvalidHeaderError(
                            f"Worksheet {spec.name!r} includes unknown columns: {unknown}"
                        )
                    selected = set(spec.include_columns) | key_headers
                else:
                    selected = available

                max_column = len(headers)
                sheet_records: list[SourceRecord] = []
                for row_number, row_cells in enumerate(
                    worksheet.iter_rows(
                        min_row=spec.header_row + 1,
                        max_col=max_column,
                    ),
                    start=spec.header_row + 1,
                ):
                    cells_read += max_column
                    if cells_read > self._max_cells:
                        raise ExcelReadError(
                            f"Workbook exceeds the configured {self._max_cells}-cell limit"
                        )
                    values: dict[str, JsonValue] = {}
                    cell_by_header: dict[str, Any] = {}
                    for index, cell in enumerate(row_cells):
                        header = headers[index]
                        if not header or header not in selected:
                            continue
                        values[header] = _json_value(cell.value)
                        cell_by_header[header] = cell
                    if not values or all(value is None or value == "" for value in values.values()):
                        continue
                    key_group = next(
                        (
                            group
                            for group in available_key_groups
                            if all(
                                values.get(column) is not None
                                and values.get(column) != ""
                                for column in group
                            )
                        ),
                        None,
                    )
                    if key_group is None:
                        raise MissingRecordKeyError(
                            f"Worksheet {spec.name!r} row {row_number} has no complete "
                            f"configured key: {available_key_groups}"
                        )
                    key = _record_key(spec.name, key_group, values)
                    refs = [
                        SourceRef(
                            drive_id=drive_id,
                            item_id=item_id,
                            version_id=version_id,
                            file_hash=file_hash,
                            sheet=spec.name,
                            row=cell.row,
                            cells=cell.coordinate,
                            raw_value=values[header],
                        )
                        for header, cell in cell_by_header.items()
                        if values[header] is not None
                    ]
                    sheet_records.append(
                        SourceRecord(
                            key=key,
                            sheet=spec.name,
                            row=row_number,
                            values=values,
                            source_refs=refs,
                        )
                    )
                identity_warnings.extend(_resolve_occurrence_keys(sheet_records))
                for record in sheet_records:
                    if record.key in keys_seen:
                        raise DuplicateRecordKeyError(
                            f"Duplicate record key {record.key!r} at "
                            f"{spec.name}!{record.row}"
                        )
                    keys_seen.add(record.key)
                records.extend(sheet_records)
        finally:
            workbook.close()

        return WorkbookSnapshot(
            drive_id=drive_id,
            item_id=item_id,
            version_id=version_id,
            file_hash=file_hash,
            captured_at=utc_now(),
            records=records,
            identity_warnings=identity_warnings,
        )


def diff_snapshots(
    baseline: WorkbookSnapshot, current: WorkbookSnapshot
) -> list[RecordChange]:
    """Compute stable-key record changes between two snapshots of the same drive item."""

    if (baseline.drive_id, baseline.item_id) != (current.drive_id, current.item_id):
        raise SnapshotIdentityMismatch("Snapshots belong to different source items")
    before_by_key = {record.key: record for record in baseline.records}
    after_by_key = {record.key: record for record in current.records}
    changes: list[RecordChange] = []
    for key in sorted(before_by_key.keys() | after_by_key.keys()):
        before = before_by_key.get(key)
        after = after_by_key.get(key)
        if before is None:
            changes.append(RecordChange(ChangeKind.CREATE, key, None, after))
        elif after is None:
            changes.append(RecordChange(ChangeKind.DELETE_CANDIDATE, key, before, None))
        elif before.content_hash != after.content_hash:
            changes.append(RecordChange(ChangeKind.UPDATE, key, before, after))
    return changes
